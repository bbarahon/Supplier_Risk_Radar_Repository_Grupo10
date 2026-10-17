#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Supplier Risk Radar - Helixa Pharmaceuticals
============================================

Lee noticias desde un .xlsx y, para cada una, decide:
  1. Is_Risk               (Y/N)  ¿es un riesgo para el suministro?
  2. Urgent                (Y/N)  ¿hay que alertar ya a compras?
  3. Affected_Supplier_IDs        proveedores afectados ("S003; S007")
y además: tipo de evento, severidad, puntuación de urgencia, acción de
mitigación (playbook) y una explicación en lenguaje natural.

Modelo híbrido y explicable:
  * Enlace de entidades: nombre del proveedor (con control de "trampas" de nombre
    parecido) o ubicación (ciudad / región / país / corredor logístico).
  * Detector de eventos por léxico de dominio (incendio, huelga, desastre natural,
    regulador, restricción de exportación, quiebra, ciberataque, escasez...).
  * Clasificador TF-IDF + regresión logística entrenado SOLO con el train,
    usado como segunda opinión cuando hay proveedor nombrado pero no se reconoce
    el evento.
  * Urgencia = severidad del evento (+ duración, - matices) + dependencia del
    proveedor (fuente única, criticidad, alternativo disponible, cobertura).
  * Memoria entre días: si el alternativo de un proveedor ya está afectado, la
    alternativa deja de proteger; las noticias de "vuelta a la normalidad"
    cierran alertas previas.

Uso:
  python radar_riesgo_proveedores.py --maestro Training.xlsx --noticias Test.xlsx
  python radar_riesgo_proveedores.py --maestro Training.xlsx --noticias Validation.xlsx
  python radar_riesgo_proveedores.py --maestro Training.xlsx --noticias LiveDay.xlsx \
         --historial Training.xlsx Validation.xlsx Test.xlsx
  (opcional) --clave AnswerKey.xlsx   puntúa con la guía oficial de los moderadores
"""
from __future__ import annotations

import argparse
import os
import re
import sys
import unicodedata
import warnings
from dataclasses import dataclass, field

import pandas as pd

warnings.filterwarnings("ignore")

# ---------------------------------------------------------------------------
# 1. Utilidades de texto
# ---------------------------------------------------------------------------

NUM_WORDS = {
    "one": 1, "two": 2, "three": 3, "four": 4, "five": 5, "six": 6, "seven": 7,
    "eight": 8, "nine": 9, "ten": 10, "eleven": 11, "twelve": 12, "a": 1,
    "several": 3, "few": 2,
}


def norm(text) -> str:
    """minúsculas, sin acentos, sin posesivos, puntuación -> espacio."""
    if text is None or (isinstance(text, float) and pd.isna(text)):
        return ""
    t = unicodedata.normalize("NFKD", str(text)).encode("ascii", "ignore").decode()
    t = t.lower().replace("’", "'")
    t = re.sub(r"'s\b", "", t)
    t = re.sub(r"[^a-z0-9]+", " ", t)
    t = re.sub(r"\s+", " ", t).strip()
    return f" {t} "


def has(text: str, pattern: str) -> bool:
    return re.search(pattern, text) is not None


def stem(w: str) -> str:
    return w[:-1] if w.endswith("s") and len(w) > 4 else w


# ---------------------------------------------------------------------------
# 2. Datos maestros: proveedores, materiales, gazetteer
# ---------------------------------------------------------------------------

LEGAL_SUFFIX = {"ltd", "ag", "co", "inc", "gmbh", "spa", "sl", "sp", "z", "o", "as", "sa"}
# Primeros tokens que por sí solos NO identifican al proveedor (lugares / palabras comunes)
GENERIC_FIRST = {
    "taizhou", "wuhan", "gujarat", "zhejiang", "rotterdam", "rhine", "boston", "darmstadt",
    "hsinchu", "nordic", "pacific", "alpine", "atlas", "andes", "benelux", "iberia",
    "celtic", "sunrise", "orion", "lumina", "stellar", "polaris", "aurum", "vectra",
    "kaveri", "kansai", "medipack", "celsius",
}
MATERIAL_STOP = {
    "type", "drug", "substance", "components", "temperature", "controlled", "road",
    "freight", "asia", "europe", "rail", "sea", "process", "quality", "control",
    "platform", "batch", "release", "testing", "laboratory", "single", "hard", "pre",
    "filled", "monoclonal", "antibody", "information", "patient", "track", "trace",
}

CITY_ALIASES = {
    "ho chi minh city": ["ho chi minh", "saigon"],
    "sao paulo": ["sao paulo"],
    "xi'an": ["xi an", "xian"],
    "bengaluru": ["bengaluru", "bangalore"],
    "lodz": ["lodz"],
    "zurich": ["zurich"],
}

# región -> ciudades de proveedores que contiene
REGIONS = {
    "gujarat": ["ahmedabad", "vadodara"],
    "zhejiang": ["taizhou", "hangzhou"],
    "telangana": ["hyderabad"],
    "karnataka": ["bengaluru"],
    "lombardy": ["milan", "bergamo"], "lombardia": ["milan", "bergamo"],
    "puerto rico": ["san juan"],
    "catalonia": ["barcelona"], "catalunya": ["barcelona"],
    "hubei": ["wuhan"],
    "kansai": ["osaka"],
    "hesse": ["darmstadt"],
    "flanders": ["antwerp"],
    "north rhine westphalia": ["cologne"],
    "rhineland palatinate": ["ludwigshafen", "mainz"],
    "rhine": ["ludwigshafen", "mainz", "cologne"],      # tramo alemán navegable
    "southern taiwan": ["kaohsiung"], "northern taiwan": ["hsinchu"],
    "massachusetts": ["boston"], "pennsylvania": ["lancaster"],
    "tamil nadu": [], "andhra pradesh": [],
}
# ciudades / puertos sin proveedor que sirven a una región con proveedores
HINTERLAND = {"ningbo": "zhejiang", "ningbo zhoushan": "zhejiang"}

COUNTRY_WORDS = {
    "India": ["india", "indian"], "China": ["china", "chinese"],
    "Switzerland": ["switzerland", "swiss"], "Ireland": ["ireland", "irish"],
    "Germany": ["germany", "german"], "Italy": ["italy", "italian"],
    "Spain": ["spain", "spanish"], "Norway": ["norway", "norwegian"],
    "Taiwan": ["taiwan", "taiwanese"], "USA": ["united states", "american"],
    "Japan": ["japan", "japanese"], "Mexico": ["mexico", "mexican"],
    "Belgium": ["belgium", "belgian"], "Puerto Rico (USA)": ["puerto rico"],
    "Netherlands": ["netherlands", "dutch"], "Singapore": ["singapore"],
    "Brazil": ["brazil", "brazilian"], "Vietnam": ["vietnam", "vietnamese"],
    "Morocco": ["morocco", "moroccan"], "Poland": ["poland", "polish"],
}

SEA_CORRIDOR = r" (red sea|suez|cape of good hope|asia to europe|asia europe|houthi|strait of malacca) "
RAIL_CORRIDOR = r" (china europe rail|eurasian rail|rail freight|trans siberian) "


@dataclass
class Supplier:
    sid: str
    name: str
    category: str
    material: str
    country: str
    city: str
    single_source: bool
    alternate: str | None
    cover_days: float | None
    criticality: str
    spend: float
    core: list = field(default_factory=list)
    mat_terms: set = field(default_factory=set)
    city_terms: list = field(default_factory=list)


def load_master(path: str):
    xl = pd.ExcelFile(path)
    sup_df = xl.parse("Suppliers")
    mat_df = xl.parse("Materials") if "Materials" in xl.sheet_names else pd.DataFrame()
    sups = {}
    for _, r in sup_df.iterrows():
        alt = r.get("Approved_Alternate_ID")
        cover = r.get("Inventory_Cover_Days")
        s = Supplier(
            sid=str(r["Supplier_ID"]).strip(), name=str(r["Supplier_Name"]).strip(),
            category=str(r["Category"]), material=str(r["Material_Supplied"]),
            country=str(r["Country"]), city=str(r["City"]),
            single_source=str(r["Single_Source"]).strip().upper() == "Y",
            alternate=None if pd.isna(alt) else str(alt).strip(),
            cover_days=None if pd.isna(cover) else float(cover),
            criticality=str(r["Criticality"]).strip(),
            spend=float(r.get("Annual_Spend_EUR", 0) or 0),
        )
        s.core = [w for w in norm(s.name).split() if w not in LEGAL_SUFFIX]
        s.mat_terms = {stem(w) for w in norm(s.material).split()
                       if len(w) >= 4 and w not in MATERIAL_STOP}
        c = norm(s.city).strip()
        s.city_terms = list({c, *[a for k, al in CITY_ALIASES.items()
                                  if norm(k).strip() == c for a in al]})
        sups[s.sid] = s
    products = {}
    for _, r in mat_df.iterrows():
        for col in ("Primary_Supplier_ID", "Alternate_Supplier_ID"):
            sid = r.get(col)
            if isinstance(sid, str):
                products.setdefault(sid, set()).update(
                    p.strip() for p in str(r["Used_In_Products"]).split(";"))
    return sups, products


# ---------------------------------------------------------------------------
# 3. Léxico de eventos (orden = prioridad; el primero que encaja define el tipo)
# ---------------------------------------------------------------------------
# (tipo, regex, severidad base, alcance)  alcance: "point" exige proveedor nombrado,
#                                          "area" admite enlace por ubicación
EVENTS = [
    ("Financial distress", r" (going concern|insolven\w*|bankrupt\w*|creditor protection|protection from its creditors|chapter 11|missed (a )?(scheduled )?(bond )?(coupon|payment)|default(ed|s)? on|receivership|administration) ", 5, "point"),
    ("Export restriction", r" (export (ban|bans|restriction\w*|licens\w*|control\w*|curb\w*|quota\w*)|ban\w* (the )?export\w*) ", 5, "area"),
    ("Regulatory action", r" (warning letter|import alert|suspend\w* (the |its )?(\w+ ){0,4}(licen[cs]e|certificate)|suspend\w* (production|manufacturing)|licen[cs]e (was |is )?suspend\w*|(gmp|glp) (certificate )?(suspen\w*|withdraw\w*)|shipments? (to \w+ )?(may be |will be |are |were )?(held|blocked|detained)|regulator\w* (halt|stop|shut)\w*|shut down by|recall ordered) ", 5, "point"),
    ("Fire / plant incident", r" (fire|blaze|explosion|explode\w*|blast|leak|collapse\w*) ", 5, "point"),
    ("Quality / contamination", r" (contaminat\w*|recall\w*|particles|sterility|out of specification|quality (issue|failure|problem)) ", 4, "point"),
    ("Cyber attack", r" (ransomware|cyber ?attack|cyberattack|hack(ed|ers)?|malware|it outage|systems? (are |were )?offline) ", 4, "point"),
    ("Natural disaster", r" (hurricane|typhoon|cyclone|earthquake|quake|flood\w*|monsoon|tsunami|wildfire\w*|landslide|tornado|volcan\w*|storm surge) ", 5, "area"),
    ("Labour strike", r" (strike|strikes|walk(ed)? out|walkout|industrial action|stoppage) ", 3, "point"),
    ("Utility outage", r" (power (rationing|cut\w*|outage\w*|shortage\w*)|blackout\w*|rationing|grid failure) ", 3, "area"),
    ("Logistics disruption", r" (low water|reroute\w*|divert\w*|port (closure|closed)|closure of .{0,30}ports?|blockade|block\w* (a |the )?(main )?(highway|road|motorway)|congestion|border closure|canal closed) ", 2, "area"),
    ("Supply shortage / delay", r" (shortage\w*|lead times?|backlog|allocation|force majeure|delay\w* shipments?|ship .{0,20}late|late deliver\w*|delays? of|(production|output) (halted|paused|stopped|suspended)|(halt\w*|paus\w*|stop\w*|suspend\w*) (production|operations|output)) ", 3, "point"),
    ("Financial distress", r" (downgrad\w*|profit warning|liquidity|debt restructuring|layoffs) ", 2, "point"),
    ("Acquisition by competitor", r" (acquir\w*|acquisition|takeover|to buy|merger|agrees to buy) ", 2, "point"),
    ("Leadership change", r" (steps down|step down|resign\w*|chief executive (has )?(left|departs?)|ceo (departure|exit)) ", 1, "point"),
]
STRIKE_AREA = r" (port|ports|terminal|dock\w*|truckers|nationwide|general strike|air traffic|rail workers) "
DISASTER_PLANT_IMPACT = r" (factor(y|ies)|plants?|production|industrial (estate|area|zone)s?|power is out|power outage|closed|close) "
REG_STOPS = r" (suspend\w*|halt\w*|held|stop\w*|ban\w*|cannot be used|shut|blocked|detained) "

RESOLUTION = r" (restored|restores|has resumed|have resumed|resumed normal|back to normal|reopened|reopens|re opened|eases|eased|lifted|protests? end\w*|cleared the|returns to normal) "
STRONG_NEGATION = r" (no supply impact|no impact on supply|no disruption to supply|supply (is|will) not (be )?affected|no industrial sites were affected|not expected to affect supply) "
SOFT_MODIFIERS = [
    (r" (warning issued|warning for|forecast\w*|expected to bring|may reach|could reach) ", "aviso/previsión, aún no ha ocurrido"),
    (r" (safety checks|for checks|inspect buildings|as a precaution|precautionary) ", "parada preventiva por inspección"),
    (r" (not affected|were not damaged|keep operating|continue operating|continues to operate) ", "la producción sigue en marcha"),
    (r" (cut|reduce|lower)\w* .{0,40}by (around |about )?([1-4]?[0-9]) percent ", "impacto parcial (<50 %)"),
    (r" minor (observations|issues|delays) ", "impacto menor"),
]
POSITIVE = r" (no observations|approv\w* (an )?expansion|receives approval|doubles|opens|opened|new (line|hall|plant|terminal)|award|wins|secures|signs|contract|record|completes) "


def duration_signal(t: str):
    """(+1 larga / -1 corta / 0, días estimados, texto)"""
    days = []
    for m in re.finditer(r" (\d+|one|two|three|four|five|six|seven|eight|nine|ten|eleven|twelve|several|few) (to (\d+|two|three|four|six) )?(hour|day|week|month)s? ", t):
        n = m.group(3) or m.group(1)
        n = int(n) if n.isdigit() else NUM_WORDS.get(n, 1)
        unit = m.group(4)
        days.append(n / 24 if unit == "hour" else n * {"day": 1, "week": 7, "month": 30}[unit])
    if has(t, r" (indefinite\w*|could take weeks|for weeks|weeks after|take weeks|several months|no date|could not say when) "):
        return 1, max(days + [30]), "duración indefinida o de semanas"
    if not days:
        return 0, None, ""
    d = max(days)
    if d >= 21:
        return 1, d, f"duración larga (~{int(d)} días)"
    if d <= 3:
        return -1, d, f"duración corta (~{max(d, 1):.0f} días)"
    return 0, d, f"duración moderada (~{int(d)} días)"


# ---------------------------------------------------------------------------
# 4. Clasificador de texto (segunda opinión, solo train)
# ---------------------------------------------------------------------------

class TextModel:
    def __init__(self):
        self.pipe = None

    def fit(self, texts, y):
        from sklearn.feature_extraction.text import TfidfVectorizer
        from sklearn.linear_model import LogisticRegression
        from sklearn.pipeline import make_pipeline
        self.pipe = make_pipeline(
            TfidfVectorizer(ngram_range=(1, 2), sublinear_tf=True, min_df=1),
            LogisticRegression(class_weight="balanced", C=5.0, max_iter=2000))
        self.pipe.fit(texts, y)
        return self

    def proba(self, text: str) -> float:
        if self.pipe is None:
            return 0.5
        return float(self.pipe.predict_proba([text])[0][list(self.pipe.classes_).index(1)])


# ---------------------------------------------------------------------------
# 5. Motor del radar
# ---------------------------------------------------------------------------

PLAYBOOK = ["Monitor", "Increase safety stock", "Shift volume to approved alternate",
            "Start dual sourcing", "Diversify geographically"]


class RiskRadar:
    URGENT_THRESHOLD = 4
    MEMORY_MIN_SEVERITY = 4   # solo eventos serios y duraderos "comprometen" a un alternativo

    def __init__(self, suppliers, products, text_model: TextModel | None = None):
        self.s = suppliers
        self.products = products
        self.tm = text_model
        self.memory = {}   # sid -> dict(date, news_id, event, severity)
        self.insights = []

    # ---------- enlace de entidades ----------
    def link_by_name(self, t: str):
        found = {}
        for s in self.s.values():
            core = s.core
            full = " " + " ".join(core) + " "
            two = " " + " ".join(core[:2]) + " " if len(core) >= 2 else None
            if full in t or (two and two in t):
                found[s.sid] = "nombrado en la noticia"
                continue
            first = core[0]
            if first not in GENERIC_FIRST and f" {first} " in t and len(first) >= 5:
                found[s.sid] = "nombrado en la noticia"
                continue
            # nombre corto + contexto del propio proveedor (material o ciudad)
            if f" {first} " in t:
                words = {stem(w) for w in t.split()}
                ctx = s.mat_terms & words   # solo material: la ciudad confunde (p.ej. "Rotterdam")
                if ctx:
                    found[s.sid] = "nombrado en la noticia (nombre + contexto)"
        return found

    def link_by_location(self, t: str, event_type: str):
        found = {}
        cities = {c for s in self.s.values() for c in s.city_terms if f" {c} " in t}
        hinter = [r for k, r in HINTERLAND.items() if f" {k} " in t]
        regions = [r for r in REGIONS if f" {r} " in t]
        target_cities = set(cities)
        for r in hinter:
            target_cities.update(REGIONS[r])
        if not cities:   # sin ciudad concreta -> se usa la región
            for r in regions:
                target_cities.update(REGIONS[r])
        for s in self.s.values():
            if any(c in target_cities for c in s.city_terms):
                found[s.sid] = f"ubicado en la zona afectada ({s.city})"
        # nivel país: restricciones de exportación o eventos de alcance nacional
        if not found and (event_type == "Export restriction" or has(t, r" (nationwide|across the country|country wide) ")):
            words = {stem(w) for w in t.split()}
            for country, kws in COUNTRY_WORDS.items():
                if any(f" {k} " in t for k in kws):
                    cands = [s for s in self.s.values() if s.country == country]
                    by_mat = [s for s in cands if s.mat_terms & words]
                    goods = [s for s in cands if s.category not in ("Services", "Logistics")]
                    for s in (by_mat or goods):
                        found[s.sid] = f"en el país afectado ({country})"
        # corredores logísticos
        if event_type in ("Logistics disruption", "Natural disaster", "Labour strike"):
            if has(t, SEA_CORRIDOR):
                for s in self.s.values():
                    if "sea freight" in s.material.lower():
                        found[s.sid] = "opera la ruta marítima afectada"
            if has(t, RAIL_CORRIDOR):
                for s in self.s.values():
                    if "rail freight" in s.material.lower():
                        found[s.sid] = "opera la ruta ferroviaria afectada"
        # filtro sectorial: "chemical plants / distributors" -> solo químicos/API/excipientes
        if has(t, r" chemical\w* ") and len(found) > 1:
            chem = {k: v for k, v in found.items()
                    if self.s[k].category in ("API", "Excipients", "Chemicals")}
            if chem:
                found = chem
        return found

    # ---------- evento y severidad ----------
    def detect_event(self, t: str):
        c = self.detect_events(t)
        return c[0] if c else (None, 0, None)

    def detect_events(self, t: str):
        """Todos los eventos que encajan, por prioridad."""
        out = []
        for etype, pat, sev, scope in EVENTS:
            if has(t, pat):
                out.append(self._refine(t, etype, sev, scope))
        return out

    def _refine(self, t, etype, sev, scope):
        """Ajusta severidad/alcance según el contexto de la noticia."""
        if etype == "Labour strike" and has(t, STRIKE_AREA):
            return "Labour strike", 2, "area"          # huelga portuaria / transporte
        if etype == "Natural disaster" and not has(t, DISASTER_PLANT_IMPACT):
            sev = 3                                    # afecta a puertos/transporte, no cierra fábricas
        if etype == "Regulatory action" and not has(t, REG_STOPS):
            sev = 3                                    # el regulador no detiene envíos/producción
        if etype == "Utility outage" and has(t, r" rationing "):
            sev = 2                                    # racionamiento = impacto parcial
        return etype, sev, scope

    # ---------- dependencia ----------
    def compromised(self, sid, date):
        m = self.memory.get(sid)
        if not m or m["severity"] < self.MEMORY_MIN_SEVERITY or not m.get("lasting"):
            return None
        if date is not None and m["date"] is not None:
            if (date - m["date"]).days > m.get("days", 30):
                return None
        return m

    def dependency(self, sid, affected, duration_days, date):
        s = self.s[sid]
        notes = []
        if s.single_source:
            adj = {"High": 1, "Medium": 0, "Low": -1}.get(s.criticality, 0)
            notes.append(f"fuente única, criticidad {s.criticality}")
        else:
            alt = s.alternate
            if alt in affected:
                adj = 0
                notes.append(f"su alternativo {alt} también está afectado por este evento")
            elif alt and (m := self.compromised(alt, date)):
                adj = 0
                notes.append(f"su alternativo {alt} ya está afectado ({m['news_id']}: {m['event']})")
            else:
                adj = -1
                notes.append(f"tiene alternativo aprobado {alt}")
            if s.criticality == "Low":
                adj -= 1
        if duration_days and s.cover_days and duration_days >= s.cover_days:
            adj += 1
            notes.append(f"la interrupción (~{int(duration_days)} d) supera su cobertura de {int(s.cover_days)} d")
        return adj, "; ".join(notes)

    # ---------- acción de mitigación ----------
    def mitigation(self, affected, urgent, sev, date):
        if not affected:
            return "No action"
        sups = [self.s[a] for a in affected]
        countries = {s.country for s in sups}
        pair_hit = any(s.alternate in affected for s in sups)
        if (len(sups) >= 2 and len(countries) == 1 and sev >= 4) or pair_hit:
            return "Diversify geographically"
        main = max(sups, key=lambda s: (s.criticality == "High", s.spend))
        if main.single_source:
            return "Start dual sourcing" if urgent else ("Increase safety stock" if sev >= 3 else "Monitor")
        if main.alternate and not self.compromised(main.alternate, date) and sev >= 4:
            return "Shift volume to approved alternate"
        if sev >= 3 and main.cover_days and main.cover_days <= 35:
            return "Increase safety stock"
        return "Monitor"

    # ---------- noticia individual ----------
    def assess(self, news_id, headline, summary, date=None):
        t = norm(f"{headline}. {summary}")
        out = dict(News_ID=news_id, Is_Risk="N", Urgent="N", Affected_Supplier_IDs="",
                   Event_Type="Not relevant", Severity=0, Urgency_Score=0,
                   Link_Method="", Mitigation_Action="No action", ML_Risk_Prob=None,
                   Duration_Days=None, Closed_Supplier_IDs="", Explanation="")
        named = self.link_by_name(t)
        cands = self.detect_events(t)
        etype, sev, scope = cands[0] if cands else (None, 0, None)
        if not named:   # sin nombre: vale el primer evento de alcance "zona"
            area = [c for c in cands if c[2] == "area"]
            if area:
                etype, sev, scope = area[0]
        ml_p = self.tm.proba(t) if self.tm else None
        out["ML_Risk_Prob"] = None if ml_p is None else round(ml_p, 2)

        # 1) noticias de resolución: cierran alertas previas
        if has(t, RESOLUTION):
            linked = named or self.link_by_location(t, "Natural disaster")
            closed = [sid for sid in linked if sid in self.memory]
            for sid in closed:
                self.memory.pop(sid, None)
            out["Event_Type"] = "Resolution / good news"
            out["Closed_Supplier_IDs"] = "; ".join(sorted(linked))
            out["Explanation"] = "No es riesgo: la noticia indica vuelta a la normalidad" + (
                f"; se cierra la alerta previa de {', '.join(closed)}." if closed else ".")
            if closed:
                self.insights.append(f"{news_id}: cierra/rebaja la alerta previa de {', '.join(closed)} (memoria entre días).")
            return out
        if has(t, STRONG_NEGATION):
            out["Explanation"] = "No es riesgo: la propia noticia indica que no hay impacto en el suministro."
            return out

        # 2) enlace a proveedores
        if named:
            linked, method = named, "nombre"
        elif etype and scope == "area":
            linked, method = self.link_by_location(t, etype), "ubicación"
        else:
            linked, method = {}, ""

        if not linked:
            if etype:
                out["Explanation"] = (f"No es riesgo: describe un evento ({etype.lower()}) pero no "
                                      "afecta a ningún proveedor de Helixa (posible trampa de nombre o lugar).")
            else:
                out["Explanation"] = "No es riesgo: sin relación con proveedores de Helixa."
            return out

        # 3) proveedor nombrado sin evento: segunda opinión del modelo de texto
        if not etype:
            if ml_p is not None and ml_p >= 0.75 and not has(t, POSITIVE):
                etype, sev, scope = "Other adverse news (ML)", 2, "point"
            else:
                out["Event_Type"] = "Positive / neutral news"
                out["Explanation"] = (f"No es riesgo: noticia positiva o neutra sobre "
                                      f"{', '.join(f'{k} {self.s[k].name}' for k in linked)}.")
                return out

        # 4) severidad
        dmod, ddays, dtxt = duration_signal(t)
        soft = [txt for pat, txt in SOFT_MODIFIERS if has(t, pat)]
        severity = max(0, min(6, sev + dmod - (1 if soft else 0)))

        # 5) urgencia por proveedor = severidad + dependencia
        affected = sorted(linked)
        best, best_note = -9, ""
        for sid in affected:
            adj, note = self.dependency(sid, affected, ddays, date)
            if method == "ubicación" and sev <= 3 and adj > 0:
                adj = 0   # impacto indirecto (puertos, transporte): la fábrica sigue produciendo
                note += "; impacto indirecto, la dependencia no eleva la urgencia"
            if severity + adj > best:
                best, best_note = severity + adj, f"{sid}: {note}"
        urgent = best >= self.URGENT_THRESHOLD

        out.update(Is_Risk="Y", Urgent="Y" if urgent else "N",
                   Affected_Supplier_IDs="; ".join(affected), Event_Type=etype,
                   Severity=severity, Urgency_Score=best, Link_Method=method,
                   Duration_Days=None if ddays is None else round(float(ddays), 1),
                   Mitigation_Action=self.mitigation(affected, urgent, severity, date))
        why = [f"Riesgo ({etype.lower()}) para " +
               ", ".join(f"{k} {self.s[k].name} [{linked[k]}]" for k in affected) + "."]
        mods = [m for m in [dtxt] + soft if m]
        why.append(f"Severidad {severity}/6" + (f" ({'; '.join(mods)})" if mods else "") +
                   f". Dependencia → {best_note}. Puntuación {best} "
                   f"{'≥' if urgent else '<'} {self.URGENT_THRESHOLD}: " +
                   ("ALERTA URGENTE." if urgent else "puede esperar al informe semanal."))
        prods = sorted({p for sid in affected for p in self.products.get(sid, [])})
        if prods:
            why.append("Productos expuestos: " + ", ".join(prods) + ".")
        out["Explanation"] = " ".join(why)

        # 6) insights + memoria
        self._insights(news_id, affected, etype, severity, date)
        for sid in affected:
            prev = self.memory.get(sid)
            if not prev or severity >= prev["severity"] or (
                    date is not None and prev["date"] is not None and date > prev["date"]):
                self.memory[sid] = dict(date=date, news_id=news_id, event=etype, severity=severity,
                                        lasting=dmod > 0, days=int(ddays or 14) if dmod <= 0 or ddays and ddays < 90 else 90)
        return out

    def _insights(self, nid, affected, etype, severity, date):
        sups = [self.s[a] for a in affected]
        if len(sups) >= 3:
            self.insights.append(f"{nid}: un solo evento ({etype}) golpea a {len(sups)} proveedores "
                                 f"({', '.join(affected)}) → riesgo de concentración geográfica.")
        for s in sups:
            if s.alternate in affected and s.sid < s.alternate:
                self.insights.append(f"{nid}: {s.sid} y su alternativo {s.alternate} están afectados a la vez "
                                     f"({s.material}); el alternativo no protege (mismo país/región).")
            prev = self.memory.get(s.sid)
            if prev and prev["news_id"] != nid:
                self.insights.append(f"{nid}: {s.sid} {s.name} acumula una segunda señal de riesgo "
                                     f"(antes {prev['news_id']}: {prev['event']})" +
                                     (" y es fuente única." if s.single_source else "."))
            if s.alternate and (m := self.compromised(s.alternate, date)) and s.alternate not in affected:
                if severity >= 3:
                    self.insights.append(f"{nid}: {s.sid} {s.name} está afectado y su alternativo {s.alternate} "
                                         f"ya lo estaba ({m['news_id']}) → {s.material} sin respaldo operativo.")
            # proveedores para los que el afectado es el alternativo
            for o in self.s.values():
                if o.alternate == s.sid and o.sid not in affected and self.compromised(o.sid, date):
                    self.insights.append(f"{nid}: {s.sid} era el respaldo de {o.sid} {o.name}, que ya estaba "
                                         f"afectado ({self.memory[o.sid]['news_id']}) → se debilita la cobertura.")

    # ---------- lote ----------
    def run(self, news: pd.DataFrame, record_insights=True):
        df = news.copy()
        df["_date"] = pd.to_datetime(df.get("Published_Date"), errors="coerce")
        df["_ord"] = range(len(df))
        rows = []
        start = len(self.insights)
        for _, r in df.sort_values(["_date", "_ord"]).iterrows():
            res = self.assess(r["News_ID"], r.get("Headline", ""), r.get("Summary", ""),
                              None if pd.isna(r["_date"]) else r["_date"])
            res["_ord"] = r["_ord"]
            rows.append(res)
        if not record_insights:
            del self.insights[start:]
        return pd.DataFrame(rows).sort_values("_ord").drop(columns="_ord").reset_index(drop=True)


# ---------------------------------------------------------------------------
# 6. Lectura de ficheros y evaluación
# ---------------------------------------------------------------------------

def find_news_sheet(path):
    xl = pd.ExcelFile(path)
    for sh in xl.sheet_names:
        d = xl.parse(sh, nrows=2)
        if "News_ID" in d.columns and "Headline" in d.columns:
            return sh, xl.parse(sh)
    raise ValueError(f"No encuentro una hoja de noticias (con News_ID y Headline) en {path}")


def has_labels(df):
    return "Is_Risk" in df.columns and df["Is_Risk"].notna().any()


def ids_set(v):
    if v is None or (isinstance(v, float) and pd.isna(v)):
        return frozenset()
    return frozenset(x.strip() for x in str(v).replace(",", ";").split(";") if x.strip())


def score(pred: pd.DataFrame, truth: pd.DataFrame):
    """Puntuación según la guía de los moderadores + métricas clásicas."""
    m = pred.merge(truth[["News_ID", "Is_Risk", "Urgent", "Affected_Supplier_IDs"]],
                   on="News_ID", suffixes=("", "_true"))
    det = urg = link = 0
    rows = []
    for _, r in m.iterrows():
        tr, pr = r["Is_Risk_true"] == "Y", r["Is_Risk"] == "Y"
        ok_det = (tr == pr) and not (not tr and r["Urgent"] == "Y")
        det += ok_det
        ok_u = ok_l = None
        if tr:
            ok_u = r["Urgent"] == r["Urgent_true"]
            ok_l = ids_set(r["Affected_Supplier_IDs"]) == ids_set(r["Affected_Supplier_IDs_true"])
            urg += ok_u
            link += ok_l
        if not ok_det or ok_u is False or ok_l is False:
            rows.append(dict(News_ID=r["News_ID"],
                             Pred=f"{r['Is_Risk']}/{r['Urgent']}/{r['Affected_Supplier_IDs'] or '-'}",
                             Real=f"{r['Is_Risk_true']}/{r['Urgent_true']}/{r['Affected_Supplier_IDs_true'] if isinstance(r['Affected_Supplier_IDs_true'], str) else '-'}"))
    n, nr = len(m), int((m["Is_Risk_true"] == "Y").sum())
    tp = int(((m.Is_Risk == "Y") & (m.Is_Risk_true == "Y")).sum())
    fp = int(((m.Is_Risk == "Y") & (m.Is_Risk_true == "N")).sum())
    fn = int(((m.Is_Risk == "N") & (m.Is_Risk_true == "Y")).sum())
    res = dict(articulos=n, deteccion=f"{det}/{n}", urgencia=f"{urg}/{nr}", enlace=f"{link}/{nr}",
               puntos=f"{det + urg + link}/{n + 2 * nr}",
               precision_riesgo=round(tp / (tp + fp), 3) if tp + fp else 0,
               recall_riesgo=round(tp / (tp + fn), 3) if tp + fn else 0)
    return res, pd.DataFrame(rows)


def load_answer_key(path):
    xl = pd.ExcelFile(path)
    frames = [xl.parse(sh) for sh in xl.sheet_names
              if "News_ID" in xl.parse(sh, nrows=1).columns and "Is_Risk" in xl.parse(sh, nrows=1).columns]
    return pd.concat(frames, ignore_index=True).drop_duplicates("News_ID")


def train_text_model(train_path):
    try:
        _, tr = find_news_sheet(train_path)
        if not has_labels(tr):
            return None
        X = [norm(f"{h}. {s}") for h, s in zip(tr.Headline, tr.Summary)]
        y = (tr.Is_Risk == "Y").astype(int).tolist()
        return TextModel().fit(X, y)
    except Exception as e:  # el radar funciona igual sin el modelo de texto
        print(f"[aviso] modelo de texto no entrenado: {e}")
        return None


# ---------------------------------------------------------------------------
# 7. Exportación
# ---------------------------------------------------------------------------

def export(news_path, sheet, news_df, pred, insights, evaluation, out_dir):
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    os.makedirs(out_dir, exist_ok=True)
    base = os.path.splitext(os.path.basename(news_path))[0]
    out_xlsx = os.path.join(out_dir, f"{base}_RESULTADOS.xlsx")
    out_csv = os.path.join(out_dir, f"{base}_respuestas.csv")

    answers = pred[["News_ID", "Is_Risk", "Urgent", "Affected_Supplier_IDs"]]
    answers.to_csv(out_csv, index=False, encoding="utf-8-sig")

    wb = load_workbook(news_path)
    ws = wb[sheet]
    header = [c.value for c in ws[1]]
    labelled = has_labels(news_df)
    cols = {}
    for name in ["Is_Risk", "Urgent", "Affected_Supplier_IDs"]:
        target = f"Pred_{name}" if labelled else name
        if target not in header:
            header.append(target)
            ws.cell(row=1, column=len(header), value=target).font = Font(bold=True)
        cols[name] = header.index(target) + 1
    by_id = pred.set_index("News_ID")
    red, amber = PatternFill("solid", fgColor="F8CBAD"), PatternFill("solid", fgColor="FFE699")
    for row in range(2, ws.max_row + 1):
        nid = ws.cell(row=row, column=header.index("News_ID") + 1).value
        if nid not in by_id.index:
            continue
        p = by_id.loc[nid]
        for name, c in cols.items():
            cell = ws.cell(row=row, column=c, value=p[name] if p[name] else None)
            if p["Urgent"] == "Y":
                cell.fill = red
            elif p["Is_Risk"] == "Y":
                cell.fill = amber

    def add_sheet(title, df, widths=None):
        if title in wb.sheetnames:
            del wb[title]
        w = wb.create_sheet(title)
        w.append(list(df.columns))
        for c in w[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="7030A0")
        for rec in df.itertuples(index=False):
            w.append([None if (isinstance(v, float) and pd.isna(v)) else v for v in rec])
        for i, col in enumerate(df.columns, 1):
            w.column_dimensions[get_column_letter(i)].width = (widths or {}).get(col, 16)
        for r in w.iter_rows(min_row=2):
            for c in r:
                c.alignment = Alignment(wrap_text=True, vertical="top")
        w.freeze_panes = "A2"
        return w

    detail = news_df[["News_ID", "Headline"]].merge(pred, on="News_ID")
    add_sheet("Radar_Detalle", detail, {"Headline": 50, "Explanation": 90,
                                        "Affected_Supplier_IDs": 22, "Event_Type": 24,
                                        "Mitigation_Action": 30})
    add_sheet("Insights", pd.DataFrame({"Insight": insights or ["(sin insights)"]}), {"Insight": 140})
    if evaluation is not None:
        res, errs = evaluation
        df = pd.DataFrame(list(res.items()), columns=["Métrica", "Valor"])
        w = add_sheet("Evaluacion", df, {"Métrica": 22, "Valor": 14})
        if len(errs):
            w.append([])
            w.append(["Errores", "Predicho (Risk/Urgent/IDs)", "Real"])
            for rec in errs.itertuples(index=False):
                w.append(list(rec))
            w.column_dimensions["B"].width = 34
            w.column_dimensions["C"].width = 34
    wb.save(out_xlsx)
    return out_xlsx, out_csv


# ---------------------------------------------------------------------------
# 8. CLI
# ---------------------------------------------------------------------------

def main(argv=None):
    ap = argparse.ArgumentParser(description="Supplier Risk Radar (Helixa)")
    ap.add_argument("--maestro", required=True, help="xlsx con hojas Suppliers y Materials (Training)")
    ap.add_argument("--noticias", required=True, nargs="+", help="xlsx de noticias a procesar")
    ap.add_argument("--historial", nargs="*", default=[],
                    help="xlsx de días anteriores para la memoria de alertas (p.ej. train, val, test)")
    ap.add_argument("--clave", help="(opcional) Answer Key de los moderadores para puntuar")
    ap.add_argument("--salida", default="resultados", help="carpeta de salida")
    ap.add_argument("--sin-ml", action="store_true", help="desactiva el clasificador de texto")
    a = ap.parse_args(argv)

    sups, prods = load_master(a.maestro)
    tm = None if a.sin_ml else train_text_model(a.maestro)
    radar = RiskRadar(sups, prods, tm)

    # memoria de días anteriores (con las PREDICCIONES del radar, no con etiquetas)
    for h in a.historial:
        _, hdf = find_news_sheet(h)
        radar.run(hdf, record_insights=False)
    if a.historial:
        print(f"Memoria cargada: {len(radar.memory)} proveedores con alertas previas.")

    key = load_answer_key(a.clave) if a.clave else None
    for path in a.noticias:
        sheet, news = find_news_sheet(path)
        radar.insights = []
        pred = radar.run(news)
        evaluation = None
        if has_labels(news):
            evaluation = score(pred, news)
            fuente = "etiquetas del propio fichero"
        elif key is not None and news.News_ID.isin(key.News_ID).any():
            evaluation = score(pred, key[key.News_ID.isin(news.News_ID)])
            fuente = "Answer Key"
        xlsx, csv = export(path, sheet, news, pred, radar.insights, evaluation, a.salida)

        print("\n" + "=" * 78)
        print(f"{os.path.basename(path)}  ({len(pred)} noticias)")
        print("=" * 78)
        show = pred[["News_ID", "Is_Risk", "Urgent", "Affected_Supplier_IDs", "Event_Type",
                     "Urgency_Score", "Mitigation_Action"]]
        with pd.option_context("display.width", 200, "display.max_colwidth", 40):
            print(show.to_string(index=False))
        print(f"\nRiesgos: {(pred.Is_Risk == 'Y').sum()}  |  Urgentes: {(pred.Urgent == 'Y').sum()}")
        if radar.insights:
            print("\nInsights:")
            for i in dict.fromkeys(radar.insights):
                print("  -", i)
        if evaluation:
            res, errs = evaluation
            print(f"\nEvaluación ({fuente}): {res}")
            if len(errs):
                print(errs.to_string(index=False))
        print(f"\n→ {xlsx}\n→ {csv}")


if __name__ == "__main__":
    main()
